"""Exports du planning pour les equipes : Excel (vacations, demi-heures, synthese) et PDF."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from planning.models import StaffingPlan

JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]
HEADER = PatternFill("solid", fgColor="1E1B4B")


def day_label(plan: StaffingPlan) -> str:
    return f"{JOURS[plan.date.weekday()]} {plan.date:%d/%m/%Y}"


def summary_rows(plan: StaffingPlan) -> list[tuple[str, str]]:
    eur = lambda v: f"{v:,.0f} €".replace(",", " ") if v is not None else "—"
    pct = lambda v: f"{100 * v:.1f} %" if v is not None else "—"
    return [
        ("Journée", day_label(plan)), ("Planning", f"{plan.get_kind_display()} (n°{plan.pk})"),
        ("Statut", plan.get_status_display()),
        ("Heures d'agents", f"{plan.agent_hours:.0f} h" if plan.agent_hours else "—"),
        ("Coût des agents", eur(plan.cost_agents)), ("Perte attendue", eur(plan.expected_loss)),
        ("VaR 95 %", eur(plan.var95)), ("Expected Shortfall 95 %", eur(plan.es95)),
        ("Service level attendu", pct(plan.expected_service_level)),
        ("Sous-capacité maximale", pct(plan.max_undercap_probability)),
    ]


def _sheet(ws, title, header, rows, widths):
    ws.title = title
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), HEADER, Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def excel_plan(plan: StaffingPlan) -> bytes:
    wb = Workbook()
    _sheet(wb.active, "Vacations", ["Début", "Fin", "Durée (h)", "Agents"],
           [(s.start.strftime("%d/%m %H:%M"), s.end.strftime("%d/%m %H:%M"), s.hours, s.agents)
            for s in plan.shifts.order_by("start", "end")], [16, 16, 12, 10])
    _sheet(wb.create_sheet(), "Demi-heures",
           ["Heure", "Besoin idéal", "Agents planifiés", "Écart", "P(sous-capacité) %", "Perte attendue €"],
           [(r.ts.strftime("%H:%M"), r.agents_required, r.agents_scheduled, r.gap,
             round(100 * (r.undercap_probability or 0), 1), round(r.expected_loss or 0, 2))
            for r in plan.intervals.order_by("ts")], [10, 14, 16, 10, 18, 16])
    _sheet(wb.create_sheet(), "Synthèse", ["Indicateur", "Valeur"], summary_rows(plan), [28, 30])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def pdf_plan(plan: StaffingPlan) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=f"Planning {day_label(plan)}")
    styles = getSampleStyleSheet()
    style = TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E1B4B")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")])])
    story = [Paragraph(f"Planning du {day_label(plan)}", styles["Title"]),
             Paragraph(f"{plan.get_kind_display()} · {plan.get_status_display()}", styles["Normal"]),
             Spacer(1, 12), Paragraph("Synthèse", styles["Heading2"]),
             Table([("Indicateur", "Valeur"), *summary_rows(plan)], colWidths=[7 * cm, 8 * cm], style=style),
             Spacer(1, 12), Paragraph("Vacations", styles["Heading2"])]
    shifts = [(s.start.strftime("%H:%M"), s.end.strftime("%H:%M (%d/%m)"), f"{s.hours:.0f} h", s.agents)
              for s in plan.shifts.order_by("start", "end")]
    if shifts:
        story.append(Table([("Début", "Fin", "Durée", "Agents"), *shifts],
                           colWidths=[3 * cm, 4.5 * cm, 3 * cm, 3 * cm], style=style))
    else:
        story.append(Paragraph("Ce planning est exprimé en effectifs par demi-heure, sans vacations.",
                               styles["Normal"]))
    doc.build(story)
    return buffer.getvalue()
